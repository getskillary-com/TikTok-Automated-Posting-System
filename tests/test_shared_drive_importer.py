from __future__ import annotations

import shutil
import unittest
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook, load_workbook

from release_task_list.task_repository import ReleaseTaskRepository
from shared.database import init_database, session
from shared_drive_importer.excel_io import INPUT_COLUMNS
from shared_drive_importer.importer import SharedDriveImporter
from video_library_system.video_repository import VideoRepository
from video_library_system.video_scanner import read_video_metadata


class SharedDriveImporterTests(unittest.TestCase):
    def setUp(self) -> None:
        self.root = Path(".test_tmp") / self.id().replace(".", "_")
        shutil.rmtree(self.root, ignore_errors=True)
        self.root.mkdir(parents=True, exist_ok=True)
        self.db_path = self.root / "group_control.db"
        self.share_root = self.root / "shared_drive"
        init_database(self.db_path)
        seed_phone(self.db_path, "PHN-MKT", "marketing")
        seed_phone(self.db_path, "PHN-SHOW", "showcase")
        self.importer = SharedDriveImporter(share_root=self.share_root, db_path=self.db_path)

    def tearDown(self) -> None:
        shutil.rmtree(self.root, ignore_errors=True)

    def test_valid_marketing_row_imports_video_caption_and_task(self) -> None:
        batch = make_batch(
            self.share_root,
            "batch-001",
            [
                {
                    "row_id": "row-1",
                    "phone_id": "PHN-MKT",
                    "video_file": "video-1.mp4",
                    "caption": "Marketing caption",
                    "scheduled_at": "2026-05-20 12:00",
                    "publish_mode": "scheduled",
                }
            ],
            {"video-1.mp4": b"video-one"},
            ready=True,
        )

        result = self.importer.scan_once(export_status=True)

        self.assertEqual(result.scanned, 1)
        self.assertEqual(result.accepted_batches, 1)
        self.assertFalse(batch.exists())
        with session(self.db_path) as connection:
            item = connection.execute("SELECT * FROM shared_import_items WHERE row_id = 'row-1'").fetchone()
            task = connection.execute("SELECT * FROM release_tasks WHERE task_id = ?", (item["task_id"],)).fetchone()
            video = connection.execute("SELECT * FROM videos WHERE video_id = ?", (item["video_id"],)).fetchone()
            video_count = connection.execute("SELECT COUNT(*) AS count FROM videos").fetchone()["count"]
            caption_count = connection.execute("SELECT COUNT(*) AS count FROM captions").fetchone()["count"]

        self.assertEqual(item["status"], "accepted")
        self.assertEqual(task["status"], "pending")
        self.assertEqual(task["phone_id"], "PHN-MKT")
        self.assertEqual(task["scheduled_at"], "2026-05-20T12:00:00-03:00")
        self.assertTrue(Path(video["file_path"]).exists())
        self.assertIn("\\accepted\\batch-001\\", str(video["file_path"]))
        self.assertEqual(video_count, 1)
        self.assertEqual(caption_count, 1)
        self.assertTrue((self.share_root / "accepted" / "batch-001").exists())
        self.assert_status_cell("batch-001", "row-1", "import_status", "accepted")

    def test_missing_ready_file_is_skipped(self) -> None:
        make_batch(
            self.share_root,
            "batch-not-ready",
            [
                {
                    "row_id": "row-1",
                    "phone_id": "PHN-MKT",
                    "video_file": "video.mp4",
                    "caption": "Caption",
                    "scheduled_at": "2026-05-20 12:00",
                    "publish_mode": "scheduled",
                }
            ],
            {"video.mp4": b"video"},
            ready=False,
        )

        result = self.importer.scan_once(export_status=True)

        self.assertEqual(result.scanned, 0)
        self.assertEqual(result.skipped, 1)
        with session(self.db_path) as connection:
            item_count = connection.execute("SELECT COUNT(*) AS count FROM shared_import_items").fetchone()["count"]
            task_count = connection.execute("SELECT COUNT(*) AS count FROM release_tasks").fetchone()["count"]
        self.assertEqual(item_count, 0)
        self.assertEqual(task_count, 0)
        self.assertTrue((self.share_root / "inbox" / "batch-not-ready").exists())

    def test_unknown_phone_rejects_row_without_task(self) -> None:
        make_batch(
            self.share_root,
            "batch-bad-phone",
            [
                {
                    "row_id": "row-bad",
                    "phone_id": "PHN-MISSING",
                    "video_file": "video.mp4",
                    "caption": "Caption",
                    "scheduled_at": "2026-05-20 12:00",
                    "publish_mode": "scheduled",
                }
            ],
            {"video.mp4": b"video"},
            ready=True,
        )

        result = self.importer.scan_once(export_status=True)

        self.assertEqual(result.rejected_batches, 1)
        with session(self.db_path) as connection:
            item = connection.execute("SELECT status, error FROM shared_import_items WHERE row_id = 'row-bad'").fetchone()
            task_count = connection.execute("SELECT COUNT(*) AS count FROM release_tasks").fetchone()["count"]
            video_count = connection.execute("SELECT COUNT(*) AS count FROM videos").fetchone()["count"]
        self.assertEqual(item["status"], "invalid")
        self.assertIn("phone_id not found", item["error"])
        self.assertEqual(task_count, 0)
        self.assertEqual(video_count, 0)
        self.assertTrue((self.share_root / "rejected" / "batch-bad-phone").exists())

    def test_account_rule_failure_does_not_create_partial_video_or_caption(self) -> None:
        make_batch(
            self.share_root,
            "batch-marketing-product",
            [
                {
                    "row_id": "row-product",
                    "phone_id": "PHN-MKT",
                    "video_file": "video.mp4",
                    "caption": "Caption",
                    "scheduled_at": "2026-05-20 12:00",
                    "publish_mode": "scheduled",
                    "product_name": "Not allowed",
                }
            ],
            {"video.mp4": b"video"},
            ready=True,
        )

        self.importer.scan_once(export_status=True)

        with session(self.db_path) as connection:
            item = connection.execute("SELECT status, error FROM shared_import_items WHERE row_id = 'row-product'").fetchone()
            task_count = connection.execute("SELECT COUNT(*) AS count FROM release_tasks").fetchone()["count"]
            video_count = connection.execute("SELECT COUNT(*) AS count FROM videos").fetchone()["count"]
            caption_count = connection.execute("SELECT COUNT(*) AS count FROM captions").fetchone()["count"]
        self.assertEqual(item["status"], "invalid")
        self.assertIn("marketing task cannot contain product fields", item["error"])
        self.assertEqual(task_count, 0)
        self.assertEqual(video_count, 0)
        self.assertEqual(caption_count, 0)

    def test_existing_video_hash_with_different_filename_is_accepted(self) -> None:
        existing = self.root / "existing.mp4"
        existing.write_bytes(b"same-content")
        VideoRepository(self.db_path).add_video(read_video_metadata(existing), title="existing")
        make_batch(
            self.share_root,
            "batch-same-content",
            [
                {
                    "row_id": "row-same-content",
                    "phone_id": "PHN-MKT",
                    "video_file": "different-name.mp4",
                    "caption": "Caption",
                    "scheduled_at": "2026-05-20 12:00",
                    "publish_mode": "scheduled",
                }
            ],
            {"different-name.mp4": b"same-content"},
            ready=True,
        )

        self.importer.scan_once(export_status=True)

        with session(self.db_path) as connection:
            item = connection.execute("SELECT status, error FROM shared_import_items WHERE row_id = 'row-same-content'").fetchone()
            task_count = connection.execute("SELECT COUNT(*) AS count FROM release_tasks").fetchone()["count"]
            video_count = connection.execute("SELECT COUNT(*) AS count FROM videos").fetchone()["count"]
        self.assertEqual(item["status"], "accepted")
        self.assertEqual(item["error"], "")
        self.assertEqual(task_count, 1)
        self.assertEqual(video_count, 2)

    def test_existing_video_filename_is_rejected_as_duplicate(self) -> None:
        existing = self.root / "duplicate.mp4"
        existing.write_bytes(b"old-content")
        VideoRepository(self.db_path).add_video(read_video_metadata(existing), title="existing")
        make_batch(
            self.share_root,
            "batch-duplicate-name",
            [
                {
                    "row_id": "row-dup-name",
                    "phone_id": "PHN-MKT",
                    "video_file": "duplicate.mp4",
                    "caption": "Caption",
                    "scheduled_at": "2026-05-20 12:00",
                    "publish_mode": "scheduled",
                }
            ],
            {"duplicate.mp4": b"new-content"},
            ready=True,
        )

        self.importer.scan_once(export_status=True)

        with session(self.db_path) as connection:
            item = connection.execute("SELECT status, error FROM shared_import_items WHERE row_id = 'row-dup-name'").fetchone()
            task_count = connection.execute("SELECT COUNT(*) AS count FROM release_tasks").fetchone()["count"]
            video_count = connection.execute("SELECT COUNT(*) AS count FROM videos").fetchone()["count"]
        self.assertEqual(item["status"], "duplicate")
        self.assertIn("duplicate video filename", item["error"])
        self.assertEqual(task_count, 0)
        self.assertEqual(video_count, 1)

    def test_export_status_reflects_release_task_updates(self) -> None:
        make_batch(
            self.share_root,
            "batch-status",
            [
                {
                    "row_id": "row-status",
                    "phone_id": "PHN-MKT",
                    "video_file": "video.mp4",
                    "caption": "Caption",
                    "scheduled_at": "2026-05-20 12:00",
                    "publish_mode": "scheduled",
                }
            ],
            {"video.mp4": b"video-status"},
            ready=True,
        )
        self.importer.scan_once(export_status=True)
        with session(self.db_path) as connection:
            task_id = connection.execute("SELECT task_id FROM shared_import_items WHERE row_id = 'row-status'").fetchone()["task_id"]

        ReleaseTaskRepository(self.db_path).update_status(task_id, "published", run_dir="run/log/path")
        self.importer.export_status(batch_id="batch-status")

        self.assert_status_cell("batch-status", "row-status", "task_status", "published")
        self.assert_status_cell("batch-status", "row-status", "run_dir", "run/log/path")

    def assert_status_cell(self, batch_id: str, row_id: str, column: str, expected: str) -> None:
        workbook = load_workbook(self.share_root / "status" / batch_id / "status.xlsx", data_only=True)
        sheet = workbook.active
        headers = [str(cell.value or "") for cell in sheet[1]]
        index = headers.index(column) + 1
        for row_number in range(2, sheet.max_row + 1):
            if str(sheet.cell(row_number, headers.index("row_id") + 1).value or "") == row_id:
                self.assertEqual(str(sheet.cell(row_number, index).value or ""), expected)
                return
        self.fail(f"row_id not found in status.xlsx: {row_id}")


def seed_phone(db_path: Path, phone_id: str, account_type: str) -> None:
    now = datetime(2026, 5, 12, 10, 0, tzinfo=timezone.utc).replace(microsecond=0).isoformat()
    with session(db_path) as connection:
        connection.execute(
            """
            INSERT INTO phones(
                phone_id, device_name, adb_serial, account_name, account_type,
                current_status, app_package, remote_video_dir, created_at, updated_at
            )
            VALUES(?, ?, ?, ?, ?, 'online_idle', 'com.zhiliaoapp.musically', '/sdcard/DCIM/Camera', ?, ?)
            """,
            (phone_id, f"{account_type} phone", f"SERIAL-{phone_id}", phone_id, account_type, now, now),
        )


def make_batch(
    share_root: Path,
    batch_id: str,
    rows: list[dict[str, str]],
    videos: dict[str, bytes],
    *,
    ready: bool,
) -> Path:
    batch_dir = share_root / "inbox" / batch_id
    batch_dir.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(INPUT_COLUMNS)
    for row in rows:
        sheet.append([row.get(column, "") for column in INPUT_COLUMNS])
    workbook.save(batch_dir / "tasks.xlsx")
    for file_name, content in videos.items():
        path = batch_dir / file_name
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_bytes(content)
    if ready:
        (batch_dir / ".ready").write_text("", encoding="utf-8")
    return batch_dir


if __name__ == "__main__":
    unittest.main()
