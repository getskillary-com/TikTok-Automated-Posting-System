from __future__ import annotations

import http.client
import json
import shutil
import threading
import unittest
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook, load_workbook

from shared.database import init_database, session
from shared_drive_importer.excel_io import INPUT_COLUMNS
from shared_drive_web.server import create_server
from shared_drive_web.static import INDEX_HTML


class SharedDriveWebTests(unittest.TestCase):
    def setUp(self) -> None:
        self.root = Path(".test_tmp") / self.id().replace(".", "_")
        shutil.rmtree(self.root, ignore_errors=True)
        self.root.mkdir(parents=True, exist_ok=True)
        self.db_path = self.root / "group_control.db"
        self.share_root = self.root / "shared_drive"
        init_database(self.db_path)
        seed_phone(self.db_path)
        self.server = create_server(
            "127.0.0.1",
            0,
            share_root=self.share_root,
            db_path=self.db_path,
            access_token="secret",
        )
        self.thread = threading.Thread(target=self.server.serve_forever, daemon=True)
        self.thread.start()
        self.port = int(self.server.server_address[1])

    def tearDown(self) -> None:
        self.server.shutdown()
        self.server.server_close()
        self.thread.join(timeout=5)
        shutil.rmtree(self.root, ignore_errors=True)

    def test_upload_endpoint_imports_batch_and_returns_status(self) -> None:
        body, content_type = multipart_body(
            fields={"batch_id": "web-test-001"},
            files=[
                ("tasks", "tasks.xlsx", xlsx_bytes(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
                ("videos", "video-1.mp4", b"video-content", "video/mp4"),
            ],
        )

        status, data = self.request_json(
            "POST",
            "/api/upload",
            body=body,
            headers={"Content-Type": content_type, "X-Upload-Token": "secret"},
        )

        self.assertEqual(status, 200)
        self.assertEqual(data["data"]["batch_id"], "web-test-001")
        self.assertEqual(data["data"]["batch"]["status"], "accepted")
        with session(self.db_path) as connection:
            task_count = connection.execute("SELECT COUNT(*) AS count FROM release_tasks").fetchone()["count"]
            item = connection.execute("SELECT status FROM shared_import_items WHERE batch_id = 'web-test-001'").fetchone()
        self.assertEqual(task_count, 1)
        self.assertEqual(item["status"], "accepted")
        self.assertTrue((self.share_root / "accepted" / "web-test-001" / "video-1.mp4").exists())
        self.assertTrue((self.share_root / "status" / "web-test-001" / "status.xlsx").exists())

    def test_upload_requires_token_when_configured(self) -> None:
        body, content_type = multipart_body(
            fields={"batch_id": "blocked"},
            files=[
                ("tasks", "tasks.xlsx", xlsx_bytes(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
                ("videos", "video-1.mp4", b"video-content", "video/mp4"),
            ],
        )

        status, data = self.request_json("POST", "/api/upload", body=body, headers={"Content-Type": content_type})

        self.assertEqual(status, 401)
        self.assertFalse(data["ok"])

    def test_accounts_page_is_separate_from_submit_page(self) -> None:
        root_status, root_html = self.request_text("GET", "/")
        account_status, account_html = self.request_text("GET", "/accounts")

        self.assertEqual(root_status, 200)
        self.assertEqual(account_status, 200)
        self.assertIn('href="/accounts"', root_html)
        self.assertNotIn('id="phoneList"', root_html)
        self.assertIn('id="phoneList"', account_html)
        self.assertIn("返回任务提交", account_html)

    def test_single_task_video_selection_survives_file_picker_cancel(self) -> None:
        self.assertIn("selectedVideoFile: null", INDEX_HTML)
        self.assertIn('const videoInput = document.getElementById("videoInput");', INDEX_HTML)
        self.assertIn('const selectedVideoName = document.getElementById("selectedVideoName");', INDEX_HTML)
        self.assertIn('videoInput.addEventListener("change"', INDEX_HTML)
        self.assertIn("} else if (state.selectedVideoFile) {", INDEX_HTML)
        self.assertIn("videoInput.required = false;", INDEX_HTML)
        self.assertIn('formData.set("video", state.selectedVideoFile, state.selectedVideoFile.name);', INDEX_HTML)
        self.assertIn("clearSelectedVideo();", INDEX_HTML)

    def test_upload_page_auto_refreshes_batch_status(self) -> None:
        self.assertIn("refreshInFlight: false", INDEX_HTML)
        self.assertIn("if (state.refreshInFlight) return;", INDEX_HTML)
        self.assertIn("setInterval(refreshAll, 5000);", INDEX_HTML)
        self.assertIn("state.batches.flatMap", INDEX_HTML)

    def test_batches_endpoint_lists_imported_batches(self) -> None:
        body, content_type = multipart_body(
            fields={"batch_id": "web-list"},
            files=[
                ("tasks", "tasks.xlsx", xlsx_bytes(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
                ("videos", "video-1.mp4", b"video-content", "video/mp4"),
            ],
        )
        self.request_json(
            "POST",
            "/api/upload",
            body=body,
            headers={"Content-Type": content_type, "X-Upload-Token": "secret"},
        )

        status, data = self.request_json("GET", "/api/batches", headers={"X-Upload-Token": "secret"})

        self.assertEqual(status, 200)
        self.assertEqual(data["data"][0]["batch_id"], "web-list")
        self.assertEqual(data["data"][0]["items"][0]["import_status"], "accepted")

    def test_batches_endpoint_reflects_live_published_task_status(self) -> None:
        body, content_type = multipart_body(
            fields={"batch_id": "web-published"},
            files=[
                ("tasks", "tasks.xlsx", xlsx_bytes(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
                ("videos", "video-1.mp4", b"video-content", "video/mp4"),
            ],
        )
        self.request_json(
            "POST",
            "/api/upload",
            body=body,
            headers={"Content-Type": content_type, "X-Upload-Token": "secret"},
        )
        with session(self.db_path) as connection:
            connection.execute(
                """
                UPDATE release_tasks
                SET status = 'published', run_dir = 'run/log/path', updated_at = '2026-05-20T13:00:00+00:00'
                """
            )

        status, data = self.request_json("GET", "/api/batches", headers={"X-Upload-Token": "secret"})

        self.assertEqual(status, 200)
        batch = data["data"][0]
        self.assertEqual(batch["batch_id"], "web-published")
        self.assertEqual(batch["latest_task_status"], "published")
        self.assertEqual(batch["task_status_summary"], {"published": 1})
        self.assertEqual(batch["items"][0]["task_status"], "published")
        self.assertEqual(batch["items"][0]["run_dir"], "run/log/path")

    def test_status_download_reflects_live_published_task_status(self) -> None:
        body, content_type = multipart_body(
            fields={"batch_id": "web-status-download"},
            files=[
                ("tasks", "tasks.xlsx", xlsx_bytes(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
                ("videos", "video-1.mp4", b"video-content", "video/mp4"),
            ],
        )
        self.request_json(
            "POST",
            "/api/upload",
            body=body,
            headers={"Content-Type": content_type, "X-Upload-Token": "secret"},
        )
        with session(self.db_path) as connection:
            connection.execute("UPDATE release_tasks SET status = 'published'")

        status, payload = self.request_bytes("GET", "/status/web-status-download/status.xlsx?token=secret")

        self.assertEqual(status, 200)
        workbook = load_workbook(BytesIO(payload), data_only=True)
        sheet = workbook.active
        headers = [cell.value for cell in sheet[1]]
        task_status_index = headers.index("task_status") + 1
        self.assertEqual(sheet.cell(row=2, column=task_status_index).value, "published")

    def test_phones_endpoint_hides_removed_phones(self) -> None:
        status, data = self.request_json("GET", "/api/phones", headers={"X-Upload-Token": "secret"})

        self.assertEqual(status, 200)
        phone_ids = {str(phone["phone_id"]) for phone in data["data"]}
        self.assertIn("PHN-WEB", phone_ids)
        self.assertNotIn("PHN-REM", phone_ids)
        web_phone = next(phone for phone in data["data"] if phone["phone_id"] == "PHN-WEB")
        self.assertIn("explorer_name", web_phone)

    def test_phone_sync_endpoint_returns_latest_phone_rows(self) -> None:
        status, data = self.request_json("POST", "/api/phones/sync", headers={"X-Upload-Token": "secret"})

        self.assertEqual(status, 200)
        self.assertEqual(data["data"]["sync"]["mode"], "disabled")
        phone_ids = {str(phone["phone_id"]) for phone in data["data"]["phones"]}
        self.assertIn("PHN-WEB", phone_ids)
        self.assertNotIn("PHN-REM", phone_ids)

    def test_phone_account_endpoint_updates_business_label_and_type(self) -> None:
        status, data = self.request_json(
            "POST",
            "/api/phones/PHN-WEB/account",
            body=json.dumps({"account_name": "账号A", "account_type": "橱窗号"}, ensure_ascii=False).encode("utf-8"),
            headers={"Content-Type": "application/json", "X-Upload-Token": "secret"},
        )

        self.assertEqual(status, 200)
        self.assertEqual(data["data"]["account_name"], "账号A")
        self.assertEqual(data["data"]["account_type"], "showcase")
        with session(self.db_path) as connection:
            phone = connection.execute("SELECT account_name, account_type FROM phones WHERE phone_id = 'PHN-WEB'").fetchone()
        self.assertEqual(phone["account_name"], "账号A")
        self.assertEqual(phone["account_type"], "showcase")

    def test_task_endpoint_imports_marketing_task(self) -> None:
        body, content_type = multipart_body(
            fields={
                "phone_id": "PHN-WEB",
                "scheduled_at": "2026-05-20T12:00",
                "caption": "Marketing caption",
                "note": "single form",
            },
            files=[("video", "marketing.mp4", b"marketing-video-content", "video/mp4")],
        )

        status, data = self.request_json(
            "POST",
            "/api/tasks",
            body=body,
            headers={"Content-Type": content_type, "X-Upload-Token": "secret"},
        )

        self.assertEqual(status, 200)
        payload = data["data"]
        self.assertEqual(payload["account_type"], "marketing")
        self.assertEqual(payload["publish_mode"], "scheduled")
        self.assertEqual(payload["batch"]["status"], "accepted")
        with session(self.db_path) as connection:
            task = connection.execute("SELECT * FROM release_tasks").fetchone()
            item = connection.execute("SELECT * FROM shared_import_items WHERE batch_id = ?", (payload["batch_id"],)).fetchone()
        self.assertEqual(task["phone_id"], "PHN-WEB")
        self.assertEqual(task["publish_mode"], "scheduled")
        self.assertEqual(task["product_search_title"], "")
        self.assertEqual(task["product_publish_name"], "")
        self.assertEqual(item["status"], "accepted")
        self.assertTrue((self.share_root / "status" / payload["batch_id"] / "status.xlsx").exists())

    def test_task_endpoint_imports_showcase_task(self) -> None:
        body, content_type = multipart_body(
            fields={
                "phone_id": "PHN-SHOW",
                "scheduled_at": "2026-05-20T13:30",
                "caption": "Showcase caption",
                "product_search_title": "search words",
                "product_publish_name": "custom product",
                "note": "showcase form",
            },
            files=[("video", "showcase.mp4", b"showcase-video-content", "video/mp4")],
        )

        status, data = self.request_json(
            "POST",
            "/api/tasks",
            body=body,
            headers={"Content-Type": content_type, "X-Upload-Token": "secret"},
        )

        self.assertEqual(status, 200)
        payload = data["data"]
        self.assertEqual(payload["account_type"], "showcase")
        self.assertEqual(payload["publish_mode"], "scheduled")
        with session(self.db_path) as connection:
            task = connection.execute("SELECT * FROM release_tasks").fetchone()
        self.assertEqual(task["phone_id"], "PHN-SHOW")
        self.assertEqual(task["publish_mode"], "scheduled")
        self.assertEqual(task["product_search_title"], "search words")
        self.assertEqual(task["product_publish_name"], "custom product")
        self.assertEqual(task["product_name"], "custom product")

    def test_showcase_task_requires_product_fields(self) -> None:
        body, content_type = multipart_body(
            fields={
                "phone_id": "PHN-SHOW",
                "scheduled_at": "2026-05-20T13:30",
                "caption": "Showcase caption",
                "product_search_title": "search words",
            },
            files=[("video", "showcase.mp4", b"showcase-missing-product-name", "video/mp4")],
        )

        status, data = self.request_json(
            "POST",
            "/api/tasks",
            body=body,
            headers={"Content-Type": content_type, "X-Upload-Token": "secret"},
        )

        self.assertEqual(status, 400)
        self.assertIn("product_publish_name", data["error"])
        self.assert_task_count(0)

    def test_marketing_task_rejects_product_fields(self) -> None:
        body, content_type = multipart_body(
            fields={
                "phone_id": "PHN-WEB",
                "scheduled_at": "2026-05-20T12:00",
                "caption": "Marketing caption",
                "product_search_title": "not allowed",
                "product_publish_name": "not allowed",
            },
            files=[("video", "marketing.mp4", b"marketing-with-product", "video/mp4")],
        )

        status, data = self.request_json(
            "POST",
            "/api/tasks",
            body=body,
            headers={"Content-Type": content_type, "X-Upload-Token": "secret"},
        )

        self.assertEqual(status, 400)
        self.assertIn("marketing task cannot contain product fields", data["error"])
        self.assert_task_count(0)

    def test_task_endpoint_allows_running_phone(self) -> None:
        body, content_type = multipart_body(
            fields={
                "phone_id": "PHN-RUN",
                "scheduled_at": "2026-05-20T12:00",
                "caption": "Queued while running",
            },
            files=[("video", "running.mp4", b"running-video-content", "video/mp4")],
        )

        status, data = self.request_json(
            "POST",
            "/api/tasks",
            body=body,
            headers={"Content-Type": content_type, "X-Upload-Token": "secret"},
        )

        self.assertEqual(status, 200)
        payload = data["data"]
        self.assertEqual(payload["account_type"], "marketing")
        with session(self.db_path) as connection:
            task = connection.execute("SELECT * FROM release_tasks").fetchone()
            phone = connection.execute("SELECT current_status FROM phones WHERE phone_id = 'PHN-RUN'").fetchone()
        self.assertEqual(task["phone_id"], "PHN-RUN")
        self.assertEqual(task["status"], "pending")
        self.assertEqual(phone["current_status"], "running")

    def test_task_endpoint_rejects_unavailable_phone(self) -> None:
        body, content_type = multipart_body(
            fields={
                "phone_id": "PHN-OFF",
                "scheduled_at": "2026-05-20T12:00",
                "caption": "Offline phone caption",
            },
            files=[("video", "offline.mp4", b"offline-video-content", "video/mp4")],
        )

        status, data = self.request_json(
            "POST",
            "/api/tasks",
            body=body,
            headers={"Content-Type": content_type, "X-Upload-Token": "secret"},
        )

        self.assertEqual(status, 400)
        self.assertIn("phone is not ready for task submission", data["error"])
        self.assert_task_count(0)

    def request_json(
        self,
        method: str,
        path: str,
        *,
        body: bytes | None = None,
        headers: dict[str, str] | None = None,
    ) -> tuple[int, dict[str, object]]:
        connection = http.client.HTTPConnection("127.0.0.1", self.port, timeout=20)
        connection.request(method, path, body=body, headers=headers or {})
        response = connection.getresponse()
        payload = response.read().decode("utf-8")
        connection.close()
        return response.status, json.loads(payload)

    def request_text(
        self,
        method: str,
        path: str,
        *,
        body: bytes | None = None,
        headers: dict[str, str] | None = None,
    ) -> tuple[int, str]:
        connection = http.client.HTTPConnection("127.0.0.1", self.port, timeout=20)
        connection.request(method, path, body=body, headers=headers or {})
        response = connection.getresponse()
        payload = response.read().decode("utf-8")
        connection.close()
        return response.status, payload

    def request_bytes(
        self,
        method: str,
        path: str,
        *,
        body: bytes | None = None,
        headers: dict[str, str] | None = None,
    ) -> tuple[int, bytes]:
        connection = http.client.HTTPConnection("127.0.0.1", self.port, timeout=20)
        connection.request(method, path, body=body, headers=headers or {})
        response = connection.getresponse()
        payload = response.read()
        connection.close()
        return response.status, payload

    def assert_task_count(self, expected: int) -> None:
        with session(self.db_path) as connection:
            task_count = connection.execute("SELECT COUNT(*) AS count FROM release_tasks").fetchone()["count"]
        self.assertEqual(task_count, expected)


def seed_phone(db_path: Path) -> None:
    now = datetime(2026, 5, 12, 10, 0, tzinfo=timezone.utc).replace(microsecond=0).isoformat()
    with session(db_path) as connection:
        connection.executemany(
            """
            INSERT INTO phones(
                phone_id, device_name, adb_serial, account_name, account_type,
                current_status, app_package, remote_video_dir, created_at, updated_at
            )
            VALUES(?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            [
                (
                    "PHN-WEB",
                    "Web Marketing Phone",
                    "SERIAL-WEB",
                    "Marketing Account",
                    "marketing",
                    "online_idle",
                    "com.zhiliaoapp.musically",
                    "/sdcard/DCIM/Camera",
                    now,
                    now,
                ),
                (
                    "PHN-SHOW",
                    "Web Showcase Phone",
                    "SERIAL-SHOW",
                    "Showcase Account",
                    "showcase",
                    "online_idle",
                    "com.zhiliaoapp.musically",
                    "/sdcard/DCIM/Camera",
                    now,
                    now,
                ),
                (
                    "PHN-RUN",
                    "Running Phone",
                    "SERIAL-RUN",
                    "Running Account",
                    "marketing",
                    "running",
                    "com.zhiliaoapp.musically",
                    "/sdcard/DCIM/Camera",
                    now,
                    now,
                ),
                (
                    "PHN-OFF",
                    "Offline Phone",
                    "SERIAL-OFF",
                    "Offline Account",
                    "marketing",
                    "offline",
                    "com.zhiliaoapp.musically",
                    "/sdcard/DCIM/Camera",
                    now,
                    now,
                ),
                (
                    "PHN-REM",
                    "Removed Phone",
                    "SERIAL-REM",
                    "Removed Account",
                    "marketing",
                    "removed",
                    "com.zhiliaoapp.musically",
                    "/sdcard/DCIM/Camera",
                    now,
                    now,
                ),
            ],
        )


def xlsx_bytes() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(INPUT_COLUMNS)
    sheet.append([
        "row-1",
        "PHN-WEB",
        "video-1.mp4",
        "Web caption",
        "2026-05-20 12:00",
        "scheduled",
        "",
        "",
        "",
        "",
        "3",
        "",
    ])
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def multipart_body(
    *,
    fields: dict[str, str],
    files: list[tuple[str, str, bytes, str]],
) -> tuple[bytes, str]:
    boundary = "----group-control-test-boundary"
    chunks: list[bytes] = []
    for name, value in fields.items():
        chunks.extend([
            f"--{boundary}\r\n".encode("utf-8"),
            f'Content-Disposition: form-data; name="{name}"\r\n\r\n'.encode("utf-8"),
            value.encode("utf-8"),
            b"\r\n",
        ])
    for field_name, filename, content, content_type in files:
        chunks.extend([
            f"--{boundary}\r\n".encode("utf-8"),
            f'Content-Disposition: form-data; name="{field_name}"; filename="{filename}"\r\n'.encode("utf-8"),
            f"Content-Type: {content_type}\r\n\r\n".encode("utf-8"),
            content,
            b"\r\n",
        ])
    chunks.append(f"--{boundary}--\r\n".encode("utf-8"))
    return b"".join(chunks), f"multipart/form-data; boundary={boundary}"


if __name__ == "__main__":
    unittest.main()
