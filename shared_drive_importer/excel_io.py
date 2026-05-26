from __future__ import annotations

from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


INPUT_COLUMNS = [
    "row_id",
    "phone_id",
    "video_file",
    "caption",
    "scheduled_at",
    "publish_mode",
    "product_link",
    "product_name",
    "product_search_title",
    "product_publish_name",
    "max_retries",
    "note",
]

REQUIRED_COLUMNS = {"row_id", "phone_id", "video_file", "caption", "scheduled_at"}

STATUS_COLUMNS = [
    "batch_id",
    "row_id",
    "phone_id",
    "video_file",
    "video_id",
    "caption_id",
    "task_id",
    "import_status",
    "task_status",
    "scheduled_at",
    "run_dir",
    "failure_reason",
    "updated_at",
]

PHONE_COLUMNS = [
    "phone_id",
    "device_name",
    "adb_serial",
    "account_name",
    "account_type",
    "current_status",
    "app_package",
    "remote_video_dir",
]


def read_task_rows(path: Path) -> list[dict[str, str]]:
    workbook = load_workbook(path, data_only=True)
    sheet = workbook.active
    headers = [
        normalize_header(cell.value)
        for cell in next(sheet.iter_rows(min_row=1, max_row=1))
    ]
    header_set = {header for header in headers if header}
    missing = sorted(REQUIRED_COLUMNS - header_set)
    if missing:
        raise ValueError("tasks.xlsx missing required columns: " + ", ".join(missing))

    rows: list[dict[str, str]] = []
    for excel_row in sheet.iter_rows(min_row=2, values_only=True):
        if not any(value not in (None, "") for value in excel_row):
            continue
        row: dict[str, str] = {}
        for index, header in enumerate(headers):
            if not header:
                continue
            value = excel_row[index] if index < len(excel_row) else ""
            row[header] = stringify_cell(value)
        rows.append(row)
    return rows


def write_task_template(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "tasks"
    sheet.append(INPUT_COLUMNS)
    sheet.append([
        "1",
        "PHN-XXXXXXXXXXXX",
        "example.mp4",
        "Have a nice day",
        "2026-05-20 18:00",
        "scheduled",
        "",
        "",
        "",
        "",
        "3",
        "Example row; delete before production use",
    ])
    format_sheet(sheet)
    save_workbook(workbook, path)


def write_phones(path: Path, phones: list[dict[str, Any]]) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "phones"
    sheet.append(PHONE_COLUMNS)
    for phone in phones:
        sheet.append([str(phone.get(column) or "") for column in PHONE_COLUMNS])
    format_sheet(sheet)
    save_workbook(workbook, path)


def write_status(path: Path, rows: list[dict[str, Any]]) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "status"
    sheet.append(STATUS_COLUMNS)
    for row in rows:
        sheet.append([str(row.get(column) or "") for column in STATUS_COLUMNS])
    format_sheet(sheet)
    save_workbook(workbook, path)


def format_sheet(sheet: Any) -> None:
    header_fill = PatternFill(fill_type="solid", fgColor="E8EEF7")
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
    sheet.freeze_panes = "A2"
    for column_cells in sheet.columns:
        letter = get_column_letter(column_cells[0].column)
        width = max(12, min(50, max(len(str(cell.value or "")) for cell in column_cells) + 2))
        sheet.column_dimensions[letter].width = width


def save_workbook(workbook: Workbook, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def normalize_header(value: object) -> str:
    return str(value or "").strip().casefold()


def stringify_cell(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.replace(microsecond=0).isoformat(sep=" ")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()
